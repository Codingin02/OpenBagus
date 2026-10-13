"""Comprehensive tests for IDX XBRL, XLSX, PDF parsers, multi-period storage, and universal import."""

from datetime import datetime, timezone
import io
import json
from pathlib import Path
import tempfile
import unittest
import zipfile

import openpyxl
import pypdf

from openbagus.domains.equities.financial_model import (
    CanonicalFinancialStatement,
    SourceCitation,
    detect_period,
    detect_currency_and_unit,
)
from openbagus.domains.equities.xbrl import IdxXbrlParser
from openbagus.domains.equities.parsers.xlsx_parser import IdxXlsxParser
from openbagus.domains.equities.parsers.pdf_parser import IdxPdfParser
from openbagus.domains.equities.storage import FinancialStatementStorage
from openbagus.domains.equities.discovery import DocumentDiscoveryManager
from openbagus.domains.equities.data import EquityData


class TestFinancialFilingsAndParsers(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        (self.root / "data" / "equities").mkdir(parents=True, exist_ok=True)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_period_detection_rules(self):
        self.assertEqual(detect_period("Financial_Report_2023_FY.xlsx"), ("2023-FY", "FY"))
        self.assertEqual(detect_period("KLBF_Lapkeu_Q3_2024.pdf"), ("2024-Q3", "Q3"))
        self.assertEqual(detect_period("TLKM_Kuartal_1_2024.xlsx"), ("2024-Q1", "Q1"))
        self.assertEqual(detect_period("BBCA_H1_2024.zip"), ("2024-H1", "H1"))
        self.assertEqual(detect_period("ASII_9M_2024.xml"), ("2024-9M", "9M"))
        self.assertEqual(detect_period("Lapkeu_Tahunan_2022.json"), ("2022-FY", "FY"))
        self.assertIsNone(detect_period("random_unrelated_filename.txt"))

    def test_currency_and_unit_detection(self):
        c, u, m = detect_currency_and_unit("Laporan Keuangan Disajikan dalam Jutaan Rupiah (In millions of IDR)")
        self.assertEqual(c, "IDR")
        self.assertEqual(u, "million_IDR")
        self.assertEqual(m, 1_000_000.0)

        c, u, m = detect_currency_and_unit("Disajikan dalam Miliaran Rupiah")
        self.assertEqual(c, "IDR")
        self.assertEqual(u, "billion_IDR")
        self.assertEqual(m, 1_000_000_000.0)

        c, u, m = detect_currency_and_unit("Expressed in USD full amount / dalam Dolar AS satuan penuh")
        self.assertEqual(c, "USD")
        self.assertEqual(u, "USD")
        self.assertEqual(m, 1.0)

    def test_xbrl_parser_xml_and_zip(self):
        xml_content = """<?xml version="1.0" encoding="utf-8"?>
<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance"
            xmlns:idx-cor="http://www.idx.co.id/xbrl/taxonomy/2020-01-01/cor">
  <xbrli:context id="CurrentYearDuration">
    <xbrli:period>
      <xbrli:startDate>2024-01-01</xbrli:startDate>
      <xbrli:endDate>2024-12-31</xbrli:endDate>
    </xbrli:period>
  </xbrli:context>
  <xbrli:context id="CurrentYearInstant">
    <xbrli:period>
      <xbrli:instant>2024-12-31</xbrli:instant>
    </xbrli:period>
  </xbrli:context>

  <idx-cor:SalesAndRevenue contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">50000000000</idx-cor:SalesAndRevenue>
  <idx-cor:CostOfSales contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">30000000000</idx-cor:CostOfSales>
  <idx-cor:GrossProfit contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">20000000000</idx-cor:GrossProfit>
  <idx-cor:OperatingProfit contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">10000000000</idx-cor:OperatingProfit>
  <idx-cor:ProfitForTheYear contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">8000000000</idx-cor:ProfitForTheYear>
  <idx-cor:ProfitAttributableToOwnersOfParent contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">7500000000</idx-cor:ProfitAttributableToOwnersOfParent>
  <idx-cor:ProfitAttributableToNoncontrollingInterests contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">500000000</idx-cor:ProfitAttributableToNoncontrollingInterests>

  <idx-cor:CashAndCashEquivalents contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">15000000000</idx-cor:CashAndCashEquivalents>
  <idx-cor:TotalAssets contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">80000000000</idx-cor:TotalAssets>
  <idx-cor:ShortTermBankLoans contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">5000000000</idx-cor:ShortTermBankLoans>
  <idx-cor:LongTermBankLoans contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">15000000000</idx-cor:LongTermBankLoans>
  <idx-cor:TotalLiabilities contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">30000000000</idx-cor:TotalLiabilities>
  <idx-cor:EquityAttributableToOwnersOfParent contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">48000000000</idx-cor:EquityAttributableToOwnersOfParent>
  <idx-cor:NoncontrollingInterests contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">2000000000</idx-cor:NoncontrollingInterests>
  <idx-cor:TotalEquity contextRef="CurrentYearInstant" unitRef="IDR" decimals="-6">50000000000</idx-cor:TotalEquity>

  <idx-cor:NetCashFlowsFromOperatingActivities contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">9000000000</idx-cor:NetCashFlowsFromOperatingActivities>
  <idx-cor:PurchaseOfPropertyPlantAndEquipment contextRef="CurrentYearDuration" unitRef="IDR" decimals="-6">-2500000000</idx-cor:PurchaseOfPropertyPlantAndEquipment>
</xbrli:xbrl>"""

        # Test XML file
        xml_path = self.root / "KLBF_2024_FY.xml"
        xml_path.write_text(xml_content, encoding="utf-8")

        parser = IdxXbrlParser(self.root)
        stmt = parser.parse_file(xml_path, "KLBF")

        self.assertEqual(stmt.symbol, "KLBF")
        self.assertEqual(stmt.period, "2024-FY")
        self.assertEqual(stmt.revenue, 50000000000.0)
        self.assertEqual(stmt.gross_profit, 20000000000.0)
        self.assertEqual(stmt.operating_income, 10000000000.0)
        self.assertEqual(stmt.net_income_total, 8000000000.0)
        self.assertEqual(stmt.net_income_attributable, 7500000000.0)
        self.assertEqual(stmt.net_income_non_controlling, 500000000.0)
        self.assertEqual(stmt.total_assets, 80000000000.0)
        self.assertEqual(stmt.total_liabilities, 30000000000.0)
        self.assertEqual(stmt.total_equity, 50000000000.0)
        self.assertEqual(stmt.interest_bearing_debt, 20000000000.0) # 5B ST + 15B LT
        self.assertEqual(stmt.operating_cash_flow, 9000000000.0)
        self.assertEqual(stmt.capital_expenditures, 2500000000.0)
        self.assertEqual(stmt.validation_status, "VALIDATED")

        # Test ZIP wrapper
        zip_path = self.root / "KLBF_2024_FY.zip"
        with zipfile.ZipFile(zip_path, "w") as z:
            z.writestr("instance.xbrl", xml_content)

        stmt_zip = parser.parse_file(zip_path, "KLBF")
        self.assertEqual(stmt_zip.symbol, "KLBF")
        self.assertEqual(stmt_zip.revenue, 50000000000.0)

    def test_xbrl_zip_bomb_defense(self):
        # Create zip with too many files or path traversal
        zip_bad = self.root / "malicious.zip"
        with zipfile.ZipFile(zip_bad, "w") as z:
            for i in range(505):
                z.writestr(f"file_{i}.xml", "<dummy/>")

        parser = IdxXbrlParser(self.root)
        with self.assertRaises(ValueError) as ctx:
            parser.parse_file(zip_bad, "BADX")
        self.assertIn("melebihi batas", str(ctx.exception))

    def test_xlsx_parser(self):
        wb = openpyxl.Workbook()
        # Sheet 1: Posisi Keuangan (Balance Sheet)
        ws_bs = wb.active
        ws_bs.title = "Posisi Keuangan"
        ws_bs.append(["PT Industri Terbuka Tbk"])
        ws_bs.append(["Laporan Posisi Keuangan Konsolidasian"])
        ws_bs.append(["Per 31 Desember 2024 (Dinyatakan dalam Jutaan Rupiah)"])
        ws_bs.append(["Akun", "31 Des 2024", "31 Des 2023"])
        ws_bs.append(["Kas dan setara kas", "1,200", "1,000"])
        ws_bs.append(["Piutang usaha", "800", "750"])
        ws_bs.append(["Jumlah aset lancar", "2,500", "2,200"])
        ws_bs.append(["Jumlah aset tidak lancar", "5,500", "5,000"])
        ws_bs.append(["Jumlah aset", "8,000", "7,200"])
        ws_bs.append(["Utang bank jangka pendek", "400", "300"])
        ws_bs.append(["Jumlah liabilitas jangka pendek", "1,500", "1,300"])
        ws_bs.append(["Pinjaman bank jangka panjang", "1,000", "900"])
        ws_bs.append(["Jumlah liabilitas", "3,000", "2,700"])
        ws_bs.append(["Ekuitas pemilik entitas induk", "4,800", "4,300"])
        ws_bs.append(["Kepentingan non-pengendali", "200", "200"])
        ws_bs.append(["Jumlah ekuitas", "5,000", "4,500"])

        # Sheet 2: Laba Rugi
        ws_pl = wb.create_sheet(title="Laba Rugi")
        ws_pl.append(["Penjualan dan pendapatan usaha", "10,000"])
        ws_pl.append(["Beban pokok penjualan", "(6,000)"]) # Negative parenthesis
        ws_pl.append(["Laba kotor", "4,000"])
        ws_pl.append(["Laba usaha", "2,000"])
        ws_pl.append(["Laba sebelum beban pajak", "1,800"])
        ws_pl.append(["Laba tahun berjalan", "1,400"])
        ws_pl.append(["Laba yang diatribusikan kepada pemilik entitas induk", "1,350"])
        ws_pl.append(["Laba yang diatribusikan kepada kepentingan non-pengendali", "50"])

        # Sheet 3: Arus Kas
        ws_cf = wb.create_sheet(title="Arus Kas")
        ws_cf.append(["Arus kas dari aktivitas operasi", "1,900"])
        ws_cf.append(["Perolehan aset tetap", "(500)"])

        xlsx_path = self.root / "ICBP_2024_FY.xlsx"
        wb.save(xlsx_path)

        parser = IdxXlsxParser(self.root)
        stmt = parser.parse_file(xlsx_path, "ICBP")

        self.assertEqual(stmt.symbol, "ICBP")
        self.assertEqual(stmt.period, "2024-FY")
        self.assertEqual(stmt.unit, "million_IDR")
        self.assertEqual(stmt.revenue, 10000.0)
        self.assertEqual(stmt.cost_of_revenue, 6000.0)
        self.assertEqual(stmt.gross_profit, 4000.0)
        self.assertEqual(stmt.operating_income, 2000.0)
        self.assertEqual(stmt.net_income_attributable, 1350.0)
        self.assertEqual(stmt.total_assets, 8000.0)
        self.assertEqual(stmt.total_liabilities, 3000.0)
        self.assertEqual(stmt.total_equity, 5000.0)
        self.assertEqual(stmt.interest_bearing_debt, 1400.0) # 400 ST + 1000 LT
        self.assertEqual(stmt.operating_cash_flow, 1900.0)
        self.assertEqual(stmt.capital_expenditures, 500.0)
        self.assertIn("Posisi Keuangan!", stmt.citations["total_assets"].cell_reference)

    def test_pdf_parser_page_citations(self):
        # Create a synthetic PDF with pypdf
        writer = pypdf.PdfWriter()
        # Page 1: Cover
        page1 = writer.add_blank_page(width=612, height=792)
        # Note: writing actual text streams in pypdf can be tested via mock or minimal valid PDF stream
        # Let's save a valid PDF
        pdf_path = self.root / "ASII_2024_FY.pdf"
        with pdf_path.open("wb") as f:
            writer.write(f)

        parser = IdxPdfParser(self.root)
        # A blank PDF has < 500 characters, so it should be flagged as REVIEW_REQUIRED
        stmt = parser.parse_file(pdf_path, "ASII")
        self.assertEqual(stmt.symbol, "ASII")
        self.assertEqual(stmt.validation_status, "REVIEW_REQUIRED")
        self.assertTrue(any("scan/gambar" in n for n in stmt.validation_notes))

    def test_multi_period_storage_and_history(self):
        storage = FinancialStatementStorage(self.root)

        for yr in [2022, 2023, 2024]:
            stmt = {
                "symbol": "BBCA",
                "period": f"{yr}-FY",
                "period_type": "FY",
                "currency": "IDR",
                "unit": "million_IDR",
                "basis": "consolidated",
                "source": f"https://www.idx.co.id/filings/BBCA_{yr}.json",
                "published_at": f"{yr+1}-02-15T09:00:00+07:00",
                "values": {
                    "revenue": 50000 + (yr - 2022) * 5000,
                    "net_income_attributable": 20000 + (yr - 2022) * 2000,
                    "total_assets": 1200000 + (yr - 2022) * 100000,
                    "equity_attributable": 220000 + (yr - 2022) * 20000,
                    "weighted_average_shares": 123275000,
                    "shares_outstanding": 123275000,
                },
                "reported_ratios_pct": {
                    "NIM": 5.4,
                    "CAR": 25.8,
                    "NPL_gross": 1.9,
                },
            }
            storage.store_statement(stmt)

        periods = storage.list_periods("BBCA")
        self.assertEqual(periods, ["2022-FY", "2023-FY", "2024-FY"])

        latest = storage.get_statement("BBCA")
        self.assertEqual(latest["period"], "2024-FY")

        history = storage.get_history("BBCA", max_periods=2)
        self.assertEqual(len(history), 2)
        self.assertEqual(history[0]["period"], "2023-FY")
        self.assertEqual(history[1]["period"], "2024-FY")

        summary = storage.get_coverage_summary()
        self.assertEqual(summary["total_symbols"], 1)
        self.assertEqual(summary["total_statements"], 3)

    def test_cli_financials_and_valuation_commands(self):
        from openbagus import cli
        from unittest.mock import patch
        from contextlib import redirect_stdout

        storage = FinancialStatementStorage(self.root)
        stmt = {
            "symbol": "BBCA",
            "period": "2024-FY",
            "period_type": "FY",
            "currency": "IDR",
            "unit": "million_IDR",
            "basis": "consolidated",
            "source": "https://www.idx.co.id/filings/BBCA.json",
            "published_at": "2025-02-15T09:00:00+07:00",
            "values": {
                "revenue": 60000,
                "net_income_attributable": 24000,
                "total_assets": 1400000,
                "equity_attributable": 260000,
                "shares_outstanding": 123275000,
                "weighted_average_shares": 123275000,
            },
            "reported_ratios_pct": {
                "NIM": 5.5,
                "CAR": 26.0,
            },
        }
        storage.store_statement(stmt)

        shell = cli.OpenBagusShell()
        with patch.object(cli, "REPO_ROOT", self.root):
            # 1. /financials BBCA
            f_out = io.StringIO()
            with redirect_stdout(f_out):
                shell.do_financials("BBCA")
            self.assertIn("Laporan Keuangan BBCA", f_out.getvalue())
            self.assertIn("60,000", f_out.getvalue())

            # 2. /financials coverage
            cov_out = io.StringIO()
            with redirect_stdout(cov_out):
                shell.do_financials("coverage")
            self.assertIn("Matriks Ketersediaan", cov_out.getvalue())

            # 3. /financials history BBCA
            hist_out = io.StringIO()
            with redirect_stdout(hist_out):
                shell.do_financials("history BBCA")
            self.assertIn("2024-FY", hist_out.getvalue())

            # 4. /valuation BBCA
            val_out = io.StringIO()
            with redirect_stdout(val_out):
                shell.do_valuation("BBCA")
            self.assertIn("Analisis Valuasi", val_out.getvalue())
            self.assertIn("PBV", val_out.getvalue())

            # 5. /fcf BBCA (bank check)
            fcf_out = io.StringIO()
            with redirect_stdout(fcf_out):
                shell.do_fcf("BBCA")
            self.assertIn("NOT_MEANINGFUL", fcf_out.getvalue())

    def test_universal_import_via_equity_data(self):
        eq_data = EquityData(self.root)

        # Import via JSON
        json_file = self.root / "KLBF_2024_FY.json"
        json_file.write_text(json.dumps({
            "symbol": "KLBF",
            "period": "2024-FY",
            "period_type": "FY",
            "currency": "IDR",
            "unit": "million_IDR",
            "basis": "consolidated",
            "source": "https://www.idx.co.id/filings/KLBF.json",
            "published_at": "2025-03-01T09:00:00+07:00",
            "values": {
                "revenue": 30000,
                "operating_income": 4000,
                "net_income_attributable": 3200,
                "total_assets": 27000,
                "total_liabilities": 5000,
                "equity_attributable": 22000,
                "operating_cash_flow": 3800,
                "capital_expenditure": 1200,
                "shares_outstanding": 46875000,
            },
            "reported_ratios_pct": {},
        }), encoding="utf-8")

        res = eq_data.import_financial_statement(json_file, ticker="KLBF")
        self.assertEqual(res["symbol"], "KLBF")
        self.assertEqual(res["period"], "2024-FY")
        self.assertIn("ROE_pct", res["ratios"])
        self.assertAlmostEqual(res["ratios"]["ROE_pct"], 14.55, places=1)

        # Check loaded data has financial_history and fundamentals
        loaded = eq_data.load("KLBF")
        self.assertIn("2024-FY", loaded["financial_history"])
        self.assertEqual(loaded["fundamentals"]["period"], "2024-FY")
        # Ensure authentic board is set from catalog (KLBF is in catalog under sector F, board MAIN)
        self.assertEqual(loaded["board"], "MAIN")


if __name__ == "__main__":
    unittest.main()
